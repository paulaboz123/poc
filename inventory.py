"""
inventory.py - builds data/manifest.csv: one row per recording, paired with its transcript
and its row in description.xlsx, with a sanity check for each row.

Recordings = video (.mp4 ...) and audio (.m4a ...) files; media_type says which. Files can
sit in any sub-folder of data/input (video/, audio/, transcript/ ...): pairing is by name.

Recording <-> transcript: same file name first, then the same name without Teams additions
(dates, "Meeting Recording"), then a close fuzzy match. pair_method says which rule matched.

Recording <-> description.xlsx (optional; columns id | name | description):
- "13. Introduction to joining data.mp4" is matched to the row with id 13 (the number at the
  start of the file name). Files without a number are matched by name.
- That row's name becomes the title; its description is kept and later used for search and
  as context for lessons. Edit names and descriptions in the Excel, not in the manifest:
  they are refreshed on every run.
- If the number matches but the names look different, the row is flagged in "check".
  A video missing from the Excel is flagged too; audio and transcripts are not.

Re-running is safe: your edits (include, crop, a corrected transcript_file) are kept, new
files are added. To redo a row from scratch, delete it from the manifest and run again.

Usage:  uv run inventory.py
"""
from __future__ import annotations

import difflib
import logging
import re
import sys
from dataclasses import dataclass
from pathlib import Path

import config
import ingest
import media
import store

log = logging.getLogger("inventory")

TEAMS_NOISE = re.compile(r"meeting recording|meeting transcript|transcript|recording", re.I)
TEAMS_STAMP = re.compile(r"\d{8}_\d{6}")                          # 20241015_100302
LEADING_NUMBER = re.compile(r"^\s*(\d+)\s*[.)_:\-]?\s*(.*)$")       # "13. Name" -> "13", "Name"


@dataclass
class Entry:
    """One row of description.xlsx."""
    id: str
    name: str
    description: str


# --- Recording <-> transcript -------------------------------------------------------
def _norm(name: str, keep_dates: bool) -> str:
    """File name as plain lowercase words, without Teams additions (and without the date
    stamp unless keep_dates)."""
    s = TEAMS_NOISE.sub(" ", Path(name).stem)
    if not keep_dates:
        s = TEAMS_STAMP.sub(" ", s)
    return " ".join(re.sub(r"[^a-z0-9]+", " ", s.lower()).split())


def clean_title(video_name: str) -> str:
    """'Weekly sync-20241015_100302-Meeting Recording' -> 'Weekly sync (2024-10-15)'."""
    stem = Path(video_name).stem
    m = TEAMS_STAMP.search(stem)
    date = f" ({m.group()[:4]}-{m.group()[4:6]}-{m.group()[6:8]})" if m else ""
    title = TEAMS_STAMP.sub(" ", TEAMS_NOISE.sub(" ", stem))
    title = " ".join(re.sub(r"[_\-]+", " ", title).split())
    return (title or stem) + date


def pair(videos: list[Path], transcripts: list[Path]) -> dict[Path, tuple[Path | None, str]]:
    """Match each recording to at most one transcript: exact name -> normalized name -> fuzzy."""
    free = list(transcripts)
    result: dict[Path, tuple[Path | None, str]] = {}
    same_name: dict[str, int] = {}   # recurring meetings share a name; then dates matter
    for v in videos:
        same_name[_norm(v.name, False)] = same_name.get(_norm(v.name, False), 0) + 1

    def take(video: Path, match: Path, how: str) -> None:
        """Record a match and remove that transcript from the free ones."""
        free.remove(match)
        result[video] = (match, how)

    for v in videos:  # 1) same file name (any folder)
        m = next((t for t in free if t.stem.lower() == v.stem.lower()), None)
        if m:
            take(v, m, "exact")
    for v in videos:  # 2) same name once Teams noise is removed, dates kept
        if v in result:
            continue
        m = next((t for t in free if _norm(t.name, True) == _norm(v.name, True)), None)
        if m:
            take(v, m, "name")
    for v in videos:  # 3) fuzzy, dates ignored; only if clearly the best candidate
        if v in result:
            continue
        scored = sorted(((difflib.SequenceMatcher(None, _norm(v.name, False), _norm(t.name, False)).ratio(), t)
                         for t in free), key=lambda x: x[0], reverse=True)
        if scored and scored[0][0] >= 0.8 and (len(scored) == 1 or scored[0][0] - scored[1][0] > 0.05):
            exact = scored[0][0] == 1.0 and same_name[_norm(v.name, False)] == 1
            take(v, scored[0][1], "name, date ignored" if exact else f"fuzzy {scored[0][0]:.2f}")
        else:
            result[v] = (None, "none")
    return result


def checks(video_min: float | None, transcript_min: float | None, how: str) -> str:
    """Sanity checks for one recording/transcript pair, as text for the 'check' column
    ('' = nothing suspicious)."""
    issues = []
    if video_min is None:
        issues.append("UNREADABLE VIDEO")
    if how == "none":
        issues.append("NO TRANSCRIPT")
    elif transcript_min is None:
        issues.append("UNREADABLE TRANSCRIPT")
    elif video_min and transcript_min > video_min + 2:
        issues.append(f"TRANSCRIPT LONGER THAN VIDEO ({transcript_min:.0f} vs {video_min:.0f} min) - wrong pair?")
    elif video_min and transcript_min < video_min * 0.5:
        issues.append(f"TRANSCRIPT COVERS ONLY {transcript_min:.0f} of {video_min:.0f} min")
    if how.startswith("fuzzy"):
        issues.append("CHECK PAIR")
    return "; ".join(issues)


# --- Recording <-> description.xlsx -----------------------------------------------------
def _simple(text: str) -> str:
    """'Intro_to Joining-Data' -> 'intro to joining data' (letters and digits only)."""
    return " ".join(re.sub(r"[\W_]+", " ", text.lower()).split())


def _norm_id(value) -> str:
    """13, 13.0, '13', '013', '13.' -> '13'. Ids that are not numbers are kept as text."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip().rstrip(".")
    return str(int(text)) if text.isdigit() else text.lower()


def _cell(row: tuple, i: int | None) -> str:
    """Text of one Excel cell ('' for empty or missing cells), spaces tidied."""
    value = row[i] if i is not None and i < len(row) else None
    return "" if value is None else " ".join(str(value).split())


def find_descriptions_file() -> Path | None:
    """DESCRIPTIONS_FILE from .env, else the first description*.xlsx anywhere in data/input."""
    if config.DESCRIPTIONS_FILE:
        path = Path(config.DESCRIPTIONS_FILE)
        return path if path.is_absolute() else config.ROOT / path
    found = sorted(p for p in config.INPUT_DIR.rglob("*.xlsx")
                   if p.stem.lower().startswith("description") and not p.name.startswith("~$"))
    if len(found) > 1:
        log.warning("Several description files, using %s (ignoring %s)",
                    found[0].name, ", ".join(p.name for p in found[1:]))
    return found[0] if found else None


def load_descriptions(path: Path) -> dict[str, Entry]:
    """Rows of the first sheet that has a header row with 'id' and 'name' (any row, any order)."""
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError("openpyxl is not installed. Run: uv add openpyxl") from e
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            cols: dict[str, int] | None = None
            for row in rows:   # find the header row
                cells = [_cell(row, i).lower() for i in range(len(row))]
                if "id" in cells and "name" in cells:
                    cols = {k: cells.index(k) for k in ("id", "name", "description") if k in cells}
                    break
            if cols is None:
                continue
            if "description" not in cols:
                log.warning("%s has no 'description' column; only names are used", path.name)
            entries: dict[str, Entry] = {}
            for row in rows:   # iteration continues after the header row
                eid = _norm_id(row[cols["id"]] if cols["id"] < len(row) else None)
                if not eid:
                    continue
                if eid in entries:
                    log.warning("%s: id %s appears twice, using the first row", path.name, eid)
                    continue
                entries[eid] = Entry(eid, _cell(row, cols["name"]), _cell(row, cols.get("description")))
            return entries
    finally:
        wb.close()
    raise ValueError(f"{path.name}: no sheet has a header row with the columns 'id' and 'name'")


def _names_agree(file_part: str, excel_name: str) -> bool:
    """True if the name in the file and the name in the Excel look like the same recording."""
    a, b = _simple(file_part), _simple(excel_name)
    return not a or not b or a in b or b in a or difflib.SequenceMatcher(None, a, b).ratio() >= 0.6


def match_entry(file_name: str, entries: dict[str, Entry]) -> tuple[Entry | None, str]:
    """The description.xlsx row for one file. Returns (row or None, problem for 'check')."""
    stem = TEAMS_STAMP.sub(" ", Path(file_name).stem)
    m = LEADING_NUMBER.match(stem)
    if m and _norm_id(m.group(1)) in entries:
        entry = entries[_norm_id(m.group(1))]
        if _names_agree(m.group(2), entry.name):
            return entry, ""
        return entry, f"EXCEL ID {entry.id} HAS ANOTHER NAME ('{entry.name}') - check"

    # no number, or a number the Excel does not have: try the name
    target = _simple(m.group(2) if m else stem)
    scored = sorted(((difflib.SequenceMatcher(None, target, _simple(e.name)).ratio(), e)
                     for e in entries.values() if e.name), key=lambda x: x[0], reverse=True)
    if scored and scored[0][0] >= 0.85 and (len(scored) == 1 or scored[0][0] - scored[1][0] > 0.05):
        entry = scored[0][1]
        if m:
            return entry, f"FILE NUMBER {m.group(1)} BUT EXCEL ID {entry.id} HAS THIS NAME - check"
        return entry, "" if scored[0][0] == 1.0 else f"MATCHED TO EXCEL ID {entry.id} BY NAME - check"
    return None, "NOT IN description.xlsx"


def _apply_entry(row: dict, entries: dict[str, Entry], used: set[str]) -> None:
    """Fill catalog_id, title and description from the Excel; add any problem to 'check'."""
    entry, problem = None, "NOT IN description.xlsx"
    for name in (row.get("video_file"), row.get("transcript_file")):
        if name:
            entry, problem = match_entry(Path(name).name, entries)
            if entry:
                break
    if entry:
        row["catalog_id"], row["description"] = entry.id, entry.description
        if entry.name:
            row["title"] = entry.name
        used.add(entry.id)
    else:   # keep anything typed into the manifest by hand
        row.setdefault("catalog_id", "")
        row.setdefault("description", "")
    if problem and (entry or row.get("media_type") == "video"):   # the Excel may list only videos
        row["check"] = "; ".join(x for x in (row.get("check"), problem) if x)


# --- Main ------------------------------------------------------------------------------
def main() -> int:
    config.setup_logging()
    if not config.INPUT_DIR.exists():
        log.error("Folder %s does not exist", config.INPUT_DIR)
        return 1

    entries: dict[str, Entry] = {}
    desc_path = find_descriptions_file()
    if desc_path is None:
        log.info("No description*.xlsx in %s: titles come from file names", config.INPUT_DIR)
    elif not desc_path.exists():
        log.error("DESCRIPTIONS_FILE %s does not exist", desc_path)
        return 1
    else:
        try:
            entries = load_descriptions(desc_path)
        except Exception as e:
            log.error("Could not read %s: %s", desc_path.name, e)
            return 1
        log.info("Read %s rows from %s", len(entries), desc_path.name)

    files = [p for p in config.INPUT_DIR.rglob("*") if p.is_file()]
    # videos first, so a transcript pairs with the video rather than with an audio copy of it
    videos = sorted((p for p in files if p.suffix.lower() in config.MEDIA_EXT),
                    key=lambda p: (p.suffix.lower() in config.AUDIO_EXT, p.as_posix().lower()))
    transcripts = sorted(p for p in files if p.suffix.lower() in config.TRANSCRIPT_EXT)
    # prefer .vtt over .docx when both exist for the same recording (timestamps are exact)
    transcripts.sort(key=lambda p: p.suffix.lower() != ".vtt")
    log.info("Found %s recordings (video/audio) and %s transcripts in %s",
             len(videos), len(transcripts), config.INPUT_DIR)

    def rel(p: Path) -> str:
        """Path relative to data/input, with forward slashes (as stored in the manifest)."""
        return p.relative_to(config.INPUT_DIR).as_posix()

    def key(r: dict) -> str:   # a row is identified by its recording, or its transcript if it has none
        """Identity of a manifest row: its recording, or its transcript if it has no recording."""
        return r.get("video_file") or "t:" + r.get("transcript_file", "")

    def new_row(file: Path, **fields) -> dict:
        """A fresh row; video_id = Excel id + name when known, so it is easy to recognise."""
        entry, _ = match_entry(file.name, entries) if entries else (None, "")
        title = entry.name if entry and entry.name else clean_title(file.name)
        label = f"{entry.id} {title}" if entry else title
        return {"video_id": store.stable_id(rel(file), label), "include": "no", "title": title,
                "crop": "", **fields}

    existing = {key(r): r for r in store.read_manifest()}
    taken = {r.get("transcript_file") for r in existing.values()}
    pairs = pair([v for v in videos if rel(v) not in existing],
                 [t for t in transcripts if rel(t) not in taken])

    video_names = {_norm(v.name, True) for v in videos if v.suffix.lower() in config.VIDEO_EXT}
    used_ids: set[str] = set()
    rows: list[dict] = []
    for v in videos:
        try:
            video_min: float | None = round(media.duration_s(v) / 60, 1)
        except media.MediaError as e:
            log.warning("%s: %s", v.name, e)
            video_min = None
        if rel(v) in existing:     # keep your edits; only refresh the measured columns
            row = existing[rel(v)]
            t_rel, how = row.get("transcript_file", ""), row.get("pair_method") or "manual"
        else:
            t_path, how = pairs[v]
            t_rel = rel(t_path) if t_path else ""
            row = new_row(v, video_file=rel(v))
        row["media_type"] = "audio" if v.suffix.lower() in config.AUDIO_EXT else "video"
        rows.append(_refresh(row, t_rel, how, video_min))
        if row["media_type"] == "audio" and _norm(v.name, True) in video_names:
            row["check"] = "; ".join(x for x in [row["check"], "AUDIO OF A VIDEO ALREADY LISTED - keep include=no"] if x)

    # Transcripts without a recording. Skip the second format of one we already have
    # (.docx next to its .vtt).
    used = {r["transcript_file"] for r in rows if r["transcript_file"]}
    used_names = {_norm(Path(u).name, True) for u in used}
    for t in transcripts:
        if rel(t) in used or _norm(t.name, True) in used_names:
            continue
        row = existing.get("t:" + rel(t)) or new_row(t, media_type="none", video_file="")
        rows.append(_refresh(row, rel(t), "transcript only", None))
        used_names.add(_norm(t.name, True))

    if entries:
        for row in rows:
            _apply_entry(row, entries, used_ids)
        missing = [e for e in entries.values() if e.id not in used_ids]
        if missing:
            log.info("%s rows of %s have no file in data/input (fine if you left them out on purpose): %s",
                     len(missing), desc_path.name, ", ".join(f"{e.id} {e.name}"[:50] for e in missing[:15]))

    store.write_manifest(rows)
    _print_summary(rows)
    flagged = sum(bool(r["check"]) for r in rows)
    log.info("Wrote %s rows to %s (%s flagged in the 'check' column)", len(rows), config.MANIFEST, flagged)
    log.info("Next: open the manifest, fix flagged rows, set include=yes for the recordings to use")
    return 0


def _refresh(row: dict, t_rel: str, how: str, video_min: float | None) -> dict:
    """Fill the measured columns: transcript length and the sanity check."""
    transcript_min = None
    if t_rel and (config.INPUT_DIR / t_rel).exists():
        try:
            cues = ingest.parse_transcript(config.INPUT_DIR / t_rel)
            transcript_min = round(cues[-1].end / 60, 1) if cues else None
        except Exception as e:
            log.warning("%s: %s", t_rel, e)
    if how == "transcript only":
        check = "" if transcript_min else "UNREADABLE TRANSCRIPT"
    else:
        check = checks(video_min, transcript_min, how if t_rel else "none")
    row.update({"transcript_file": t_rel, "pair_method": how,
                "video_min": "" if video_min is None else video_min,
                "transcript_min": "" if transcript_min is None else transcript_min,
                "check": check})
    return row


def _print_summary(rows: list[dict]) -> None:
    """One line per row, so you can see the result without opening the CSV."""
    print(f"\n{'include':7}  {'type':5}  {'id':>4}  {'title':40}  {'min':>5}  {'transcript':10}  check")
    for r in sorted(rows, key=lambda r: (r.get("media_type") != "video", r.get("media_type", ""), r["title"].lower())):
        transcript = Path(r["transcript_file"]).suffix if r["transcript_file"] else "-"
        print(f"{r['include']:7}  {r.get('media_type', ''):5}  {r.get('catalog_id', ''):>4}  "
              f"{r['title'][:40]:40}  {str(r['video_min'] or '-'):>5}  {transcript:10}  {r['check'] or 'OK'}")
    print()


if __name__ == "__main__":
    sys.exit(main())
